#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""读一份平台导出的总表（xlsx）-> 内存行对象数组。

**这是 MIGRATION.md R9 的 `parse*` 那一半。** R9 要求每份总表对应一对函数：一个
`fetch*` 把表下到磁盘（在 mssw_client.js），一个 `parse*` 把它读成内存行对象。
本仓库的 Excel 读写全是 openpyxl，JS 侧读不了 xlsx，所以 `parse*` 只能落在 Python
这边，由 mssw_parser.js 拉起。

本文件是**通用**的读表器，不是告警表专用：将来事件表/漏洞表要用，换个 `columns`
白名单就够了，不要再写第二个读表器（那正是 R9 想避免的事）。

行对象的形态（stats/metrics/*.js 消费的就是这个）：

    { "告警定性": "未知威胁" }

四条硬规矩，违反哪一条都会让取值层静默算错：

  1. **键是表头文字，绝不按列下标。** 平台导出的列集是数据相关的：同一个接口同一组
     参数，116 条数据出 77 列、0 条数据只出 68 列，整列无值的列直接不出现
     （见 mssw_client.js 里 MSSW_ALERT_TABLE_FIELDS 的注记）。按列下标写死的代码，
     换一份数据就会静默数到别的列上。
  2. **每一行都携带请求的每一列的键**（该列在表头里存在时），空值写空串 ''。
     取值层用 `hasOwnProperty` 判「这一列在不在」；漏键会把「列存在但整列为空」
     误判成「列不存在」，于是本该记「空表」的那格变成报「列名被改了」。
  3. **请求的列在表头里不存在时，一行都不带这个键**（不补空串）—— 这正是规矩 2 里
     那条防线要的信号：`hasOwnProperty` 为假 => 列名对不上 => 取值层抛错（R10）。
  4. **表头重名直接抛错。** 行对象以表头文字为键，重名列会互相覆盖，静默丢掉一列
     比报错糟得多；这里不做「取第几个」的猜测。

单元格值的形态见 `format_cell_value`（时间列给的是 **epoch 秒**，R11）。

用法:
    python parse_table_rows.py '<payload_json>'    # 建议过一遍 B64: 前缀（中文路径）
payload_json:
    {
      "path":    "<xlsx 绝对路径>",
      "sheet":   null,              # null / 省略 = 第一个 worksheet
      "columns": ["告警定性"]        # 省略 = 全部具名列（小心：大表会撞 maxBuffer）
    }

输出:
    stdout: {"sheet", "header", "rowCount", "rows", "missingColumns", "unnamedColumns"}
        header          该表的**全部**表头（诊断用；rowCount/rows 只带 columns 要的列）
        rowCount        数据行数（不含表头行）
        rows            行对象数组，见上面四条规矩
        missingColumns  请求了但表头里没有的列
        unnamedColumns  表头为空的列数（这些列无法按名字寻址，不进行对象）
    stderr: 失败时 {"diag":"error","message"}（与 build_report.py 同一套协议）
"""

import datetime
import json
import sys
from pathlib import Path

# 集群上 skill 根是只读挂载：绝不让 Python 往那儿写 __pycache__。
# 放在下面这些本地 import 之前，否则它们自己就已经先把 pycache 写进去了。
sys.dont_write_bytecode = True

# 报错文案里会有「」和中文列名；PYTHONIOENCODING 由调用方的 workdir.childEnv() 设好，
# 但手工在 GBK 控制台执行时 stdout 可能是 GBK，那样会在**已经读完表**之后抛
# UnicodeEncodeError。只放宽错误处理、不改编码，与 build_report.py 同款。
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, 'reconfigure'):
        _stream.reconfigure(errors='replace')

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook

from _path_helper import decode_argv
# 行集合的口径必须与合并引擎同源，否则取值层数出的行数与交付物里看到的对不上
from build_report import (
    collect_source_rows,
    is_blank,
    js_string,
    load_first_worksheet,
    split_leading_blank_rows,
)


def load_source_sheet(file_path, sheet_name=None):
    """读一份导出的总表 -> (实际 sheet 名, worksheet)。

    sheet_name 为空时取第一个 worksheet，走的是合并引擎那条路
    （load_first_worksheet）：**普通模式，不是 read_only**。平台有
    `<dimension ref="A1"/>` 声明不准的导出，read_only 会只读到 A1（见
    _path_helper.reset_read_only_dimensions 的说明）。data_only=True 与引擎一致：
    要的是缓存值，不是公式。
    """
    if not file_path:
        raise ValueError('总表路径为空')
    path = Path(file_path)
    if not path.exists():
        raise ValueError(f'总表不存在: {file_path}')

    if not sheet_name:
        _workbook, name, sheet = load_first_worksheet(file_path)
        return name, sheet

    workbook = load_workbook(path, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f'总表里没有 sheet 「{sheet_name}」: {file_path}'
            f'（现有: {"、".join(workbook.sheetnames)}）'
        )
    return sheet_name, workbook[sheet_name]


def format_cell_value(value):
    """单元格值 -> 交给 JS 的形态。

    · None  -> ''（空串，不是 null：取值层按「空串 = 有这一列但没值」处理）
    · 时间  -> **epoch 秒**（R11：全仓时间戳单位统一用秒）

    时间那一条是**刻意与交付物不同**的：报告 sheet 里时间列写的是 Excel 序列号 +
    日期格式（见 build_report.to_serial_number，那是为了与旧版逐格对齐）。读到这里
    已经是 openpyxl 还原出的 datetime，把它换成秒，取值层就只有一个时间单位可用，
    不必每次去猜拿到的是天还是秒。

    无时区的 datetime 按 **UTC** 解释：Excel 里存的是墙上时间、本来就没有时区信息，
    而本地（Windows）与集群（Linux）的本地时区可能不同 —— 选 UTC 是为了同一份表
    在两处读出同一个数。time 类型（纯时刻，没有日期）给的是**当日零点起的秒数**，
    它不是一个时间戳，套 epoch 只会得到一个荒唐的 1900 年的数。
    """
    if value is None:
        return ''
    if isinstance(value, bool):
        return value
    if isinstance(value, datetime.datetime):
        if value.tzinfo is None:
            value = value.replace(tzinfo=datetime.timezone.utc)
        return int(value.timestamp())
    if isinstance(value, datetime.date):
        return int(datetime.datetime(
            value.year, value.month, value.day, tzinfo=datetime.timezone.utc
        ).timestamp())
    if isinstance(value, datetime.time):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, (int, float)):
        return value
    return str(value)


def read_header_and_rows(sheet):
    """sheet -> (header, 数据行值列表)。

    行集合与报告里这张 sheet **完全同源**：同一个 collect_source_rows（整行没值才丢，
    空串算有内容）+ 同一个 split_leading_blank_rows（丢掉开头的空行）。两处同源，
    取值层数出来的行数才等于交付物里肉眼看到的行数。
    """
    rows = split_leading_blank_rows(collect_source_rows(sheet))
    if not rows:
        return [], []

    # 表头文字原样用，**不去首尾空白**：键就是要与平台导出的列名逐字对上，
    # 名字里真带了空格的话，宁可让 hasOwnProperty 那条防线抛错（错误信息里带着
    # 完整表头，一眼看得出来），也不要自作主张替用户决定「空白算不算」
    header = ['' if is_blank(value) else js_string(value) for value in rows[0][0]]
    return header, [values for values, _formats in rows[1:]]


def parse_table_rows(payload):
    """R9 的 parse*：总表 -> 行对象数组。返回体见模块头注。"""
    sheet_name, sheet = load_source_sheet(payload.get('path'), payload.get('sheet') or None)

    requested = payload.get('columns')
    if requested is not None:
        if not isinstance(requested, list) or not all(isinstance(c, str) for c in requested):
            raise ValueError('columns 必须是字符串数组（或省略表示要全部具名列）')
        requested = [column for column in requested if column]

    header, value_rows = read_header_and_rows(sheet)
    named = [name for name in header if name]

    duplicated = sorted({name for name in named if named.count(name) > 1})
    if duplicated:
        raise ValueError(
            f'总表 {sheet_name} 的表头有重名列: {"、".join(duplicated)}。'
            '行对象以表头文字为键，重名列会互相覆盖 —— 这里不做「取第几个」的猜测，'
            '先让平台把列名定下来。'
        )

    # 列名 -> 列序（0 基）。named 已保证无不重名，取到的就是唯一那一列。
    index_of = {}
    for index, name in enumerate(header):
        if name and name not in index_of:
            index_of[name] = index

    wanted = named if requested is None else [column for column in requested if column]
    # 表头里没有的列不进行对象（规矩 3），只回报给调用方
    missing = [column for column in wanted if column not in index_of]
    keys = [column for column in wanted if column in index_of]

    data = []
    for values in value_rows:
        record = {}
        for key in keys:
            index = index_of[key]
            record[key] = format_cell_value(values[index] if index < len(values) else None)
        data.append(record)

    return {
        'sheet': sheet_name,
        'header': header,
        'rowCount': len(data),
        'rows': data,
        'missingColumns': missing,
        'unnamedColumns': len(header) - len(named),
    }


def main():
    decode_argv()
    if len(sys.argv) < 2:
        raise SystemExit("Usage: parse_table_rows.py '<payload_json>'")

    try:
        payload = json.loads(sys.argv[1])
        if not isinstance(payload, dict):
            raise ValueError('payload 必须是一个 JSON 对象')
        result = parse_table_rows(payload)
    except Exception as error:  # noqa: BLE001 —— 统一转成可读 JSON 交给 JS 侧
        print(json.dumps({
            'diag': 'error',
            'message': str(error),
        }, ensure_ascii=False), file=sys.stderr)
        return 1

    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
