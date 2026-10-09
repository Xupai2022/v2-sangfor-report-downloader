#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
资产表加工脚本。

与老仓库 health-checkup-report/scripts/process_risk_list_table.py 的 asset 分支相比，
**唯一差异是去掉了「删除 5 列」那一步**：

    老仓库删列清单（本脚本不执行）：
        zdy / 责任人电话 / 责任人(设备上报) / 实时认证用户名 / 托管状态
    即删除 user_info 整组 + mss_service_info 整组 + custom_attribute 组。

保留的两步：
    1) 清洗「所属业务」列：按逗号分隔，每段去掉 '/' 及其前面的部分
    2) 合并待审核资产，并新增「审核状态」列（台账填已审核、待审核表填待审核）

用法:
    python process_asset_table.py <台账.xlsx> <output_dir> [待审核.xlsx]

输出:
    stdout: {"filePath": "<output_dir>/<原文件名>.xlsx"}
"""
import json
import os
import sys

# 集群上 skill 根是只读挂载：绝不让 Python 往那儿写 __pycache__（须在本地 import 之前）
sys.dont_write_bytecode = True

from openpyxl import load_workbook
from openpyxl.styles import Font as _Font, PatternFill

from _path_helper import decode_argv
decode_argv()

BUSINESS_COLUMN_ALIASES = ['所属业务']
APPROVAL_COLUMN_NAME = '审核状态'


def normalize(value):
    return '' if value is None else str(value).strip()


def clean_business_column(value):
    """处理所属业务列：按逗号分隔，每个分段去掉 '/' 及其前面的部分，保留后面的内容。"""
    text = normalize(value)
    if not text:
        return text
    parts = text.split(',')
    cleaned = []
    for part in parts:
        part = part.strip()
        if '/' in part:
            part = part.rsplit('/', 1)[-1].strip()
        if part:
            cleaned.append(part)
    return ', '.join(cleaned) if cleaned else ''


def build_col_header(ws):
    """查找第一个非空行作为表头，返回 (col_name->index 字典, header_row_number)。"""
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        values = [normalize(cell) for cell in row]
        if any(values):
            return {name: i for i, name in enumerate(values) if name}, idx
    return {}, 1


def find_business_col_idx(col_map):
    for alias in BUSINESS_COLUMN_ALIASES:
        if alias in col_map:
            return col_map[alias]
    return None


def clean_business_values(ws, header_row, col_idx):
    """清洗所属业务列，只写回真正发生变化的值。"""
    if col_idx is None:
        return
    for row in ws.iter_rows(min_row=header_row + 1):
        if col_idx >= len(row):
            break
        cell = row[col_idx]
        if cell.value is None:
            continue
        raw = normalize(cell.value)
        cleaned = clean_business_column(raw)
        if cleaned != raw:
            cell.value = cleaned


def append_approval_column(ws, header_row, value):
    """在表头行末尾追加「审核状态」列，并给现有数据行填 value。返回新列号（1-indexed）。"""
    approval_col_idx = (ws.max_column or 0) + 1
    header_cell = ws.cell(row=header_row, column=approval_col_idx)
    header_cell.value = APPROVAL_COLUMN_NAME

    # 表头样式对齐 MSSW 导出风格：底色 #333333，字体从第一列表头复制但改白
    try:
        header_cell.fill = PatternFill(start_color='FF333333', end_color='FF333333', fill_type='solid')
        first_header_cell = ws.cell(row=header_row, column=1)
        if first_header_cell.has_style:
            base_font = first_header_cell.font
            header_cell.font = _Font(
                name=base_font.name or '微软雅黑',
                size=base_font.size,
                bold=base_font.bold,
                italic=base_font.italic,
                underline=base_font.underline,
                strike=base_font.strike,
                vertAlign=base_font.vertAlign,
                color='FFFFFFFF'
            )
            header_cell.border = first_header_cell.border.copy()
            header_cell.alignment = first_header_cell.alignment.copy()
            header_cell.number_format = first_header_cell.number_format
            header_cell.protection = first_header_cell.protection.copy()
    except Exception as e:
        sys.stderr.write(f'[WARN] 设置审核状态表头样式失败: {e}\n')

    for row in ws.iter_rows(min_row=header_row + 1):
        ws.cell(row=row[0].row, column=approval_col_idx, value=value)

    return approval_col_idx


def merge_wait_approve(ws, header_row, approval_col_idx, wait_approve_path):
    """把待审核表的数据行追加到台账末尾，填「待审核」。

    列数超出台账的列直接丢弃（与老仓库一致）。
    """
    wb_wait = load_workbook(wait_approve_path)
    ws_wait = wb_wait.active
    wait_col_map, wait_header_row = build_col_header(ws_wait)
    wait_business_col_idx = find_business_col_idx(wait_col_map)

    current_max_row = ws.max_row or header_row
    target_col_count = approval_col_idx - 1  # 不含审核状态列

    for row_idx, row in enumerate(
        ws_wait.iter_rows(min_row=wait_header_row + 1, values_only=True),
        start=current_max_row + 1
    ):
        if not any(normalize(c) for c in row):
            continue
        for c_idx, cell_value in enumerate(row, start=1):
            if c_idx > target_col_count:
                break
            ws.cell(row=row_idx, column=c_idx, value=cell_value)

        if wait_business_col_idx is not None and wait_business_col_idx < target_col_count:
            cell = ws.cell(row=row_idx, column=wait_business_col_idx + 1)
            if cell.value is not None:
                raw = normalize(cell.value)
                cleaned = clean_business_column(raw)
                if cleaned != raw:
                    cell.value = cleaned

        ws.cell(row=row_idx, column=approval_col_idx, value='待审核')

    wb_wait.close()


def main():
    if len(sys.argv) < 3:
        raise SystemExit('Usage: process_asset_table.py <asset.xlsx> <output_dir> [wait_approve.xlsx]')

    input_path = sys.argv[1]
    output_dir = sys.argv[2]
    wait_approve_path = sys.argv[3] if len(sys.argv) >= 4 else None

    if not os.path.isfile(input_path):
        raise SystemExit(f'输入文件不存在: {input_path}')

    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, os.path.basename(input_path))

    wb = load_workbook(input_path)
    ws = wb.active
    col_map, header_row = build_col_header(ws)

    # 1) 清洗「所属业务」列
    clean_business_values(ws, header_row, find_business_col_idx(col_map))

    # 2) 追加审核状态列 + 合并待审核数据
    approval_col_idx = append_approval_column(ws, header_row, '已审核')
    if wait_approve_path and os.path.isfile(wait_approve_path):
        try:
            merge_wait_approve(ws, header_row, approval_col_idx, wait_approve_path)
        except Exception as e:
            # 合并失败不阻断主流程，但记录到 stderr
            sys.stderr.write(f'[WARN] 合并待审核资产失败: {e}\n')

    wb.save(output_path)
    wb.close()
    print(json.dumps({'filePath': output_path}, ensure_ascii=False))


if __name__ == '__main__':
    main()
